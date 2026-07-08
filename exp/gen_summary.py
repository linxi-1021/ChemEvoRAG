"""Generate final summary Excel from all baseline results."""
import json,glob,os
from collections import defaultdict
import openpyxl

dir='D:/Desktop/evo/ChemEvoRAG_Phase1-main/data/baseline_results'
files={}
def avg_score(path):
    with open(path,encoding='utf-8') as fp:
        results=json.load(fp).get('results',[])
    scores=[r.get('judge_score',0) for r in results]
    return len(results), sum(scores)/len(scores) if scores else 0, path

for f in glob.glob(f'{dir}/*.json'):
    name=os.path.basename(f).split('_20')[0]
    nq,avg,_=avg_score(f)
    if name not in files:
        files[name]=(f,nq,avg)
    else:
        _,prev_nq,prev_avg=files[name]
        # Pick file with most questions; tie-break by highest avg score
        if nq>prev_nq or (nq==prev_nq and avg>prev_avg):
            files[name]=(f,nq,avg)
files={k:v[0] for k,v in files.items()}
files['ChemEvoRAG']='D:/Desktop/evo/ChemEvoRAG_Phase1-main/data/eval/eval_results.json'

wb=openpyxl.Workbook()
ws=wb.active; ws.title='Summary'
intents=['alias_resolution','entity_lookup','property_query','reaction_comparison','reaction_condition_query']
ws.cell(row=1,column=1,value='Baseline')
for c,i in enumerate(intents,2): ws.cell(row=1,column=c,value=i[:10])
ws.cell(row=1,column=7,value='OVERALL')

print(f'{"Baseline":15s}  {"ALIAS":>7s}  {"ENTITY":>7s}  {"PROP":>7s}  {"COMP":>7s}  {"COND":>7s}  {"OVERALL":>7s}')
print('-'*70)
for r_idx,(name,fpath) in enumerate(sorted(files.items()),2):
    with open(fpath,encoding='utf-8') as fp:
        data=json.load(fp)
        results=data['results'] if isinstance(data,dict) else data
    ws.cell(row=r_idx,column=1,value=name)
    bi=defaultdict(list)
    for r in results: bi[r.get('intent','?')].append(r.get('judge_score',0))
    row=f'{name:15s}'
    for c,i in enumerate(intents,2):
        s=bi[i]; ws.cell(row=r_idx,column=c,value=round(sum(s)/len(s),3) if s else 0)
        row+=f'  {sum(s)/len(s):7.3f}' if s else '     N/A'
    all_s=[r.get('judge_score',0) for r in results]
    ws.cell(row=r_idx,column=7,value=round(sum(all_s)/len(all_s),3))
    row+=f'  {sum(all_s)/len(all_s):7.3f}'
    print(row)
print('-'*70)
wb.save(f'{dir}/summary_final.xlsx')
print(f'Saved: {dir}/summary_final.xlsx')
