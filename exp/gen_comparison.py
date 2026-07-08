"""Generate comparison_final.xlsx — Detailed QA + Summary sheets for all baselines.
Matches the format of the existing comparison_final.xlsx exactly."""
import json, os, glob, sys, io
from collections import defaultdict
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

DIR = 'D:/Desktop/evo/ChemEvoRAG_Phase1-main/data/baseline_results'
CHEMEVORAG = 'D:/Desktop/evo/ChemEvoRAG_Phase1-main/data/eval/eval_results.json'

INTENTS = ['alias_resolution', 'entity_lookup', 'property_query',
           'reaction_comparison', 'reaction_condition_query']
INTENT_SHORT = {'alias_resolution': 'ALIAS', 'entity_lookup': 'ENTITY',
                'property_query': 'PROP', 'reaction_comparison': 'COMP',
                'reaction_condition_query': 'COND'}

# ---- Find best file per baseline ----
def best_file(pattern):
    best = None
    for f in glob.glob(f'{DIR}/{pattern}_*.json'):
        with open(f, encoding='utf-8') as fp:
            data = json.load(fp)
        results = data.get('results', [])
        scores = [r.get('judge_score', 0) for r in results if 'judge_score' in r]
        nq = len(results)
        avg = sum(scores) / len(scores) if scores else 0
        if best is None or nq > best[1] or (nq == best[1] and avg > best[2]):
            best = (f, nq, avg)
    return best[0] if best else None

def load_results(path):
    with open(path, encoding='utf-8') as fp:
        data = json.load(fp)
    return data.get('results', data if isinstance(data, list) else [])

# Baseline display names and files
BASELINES = [
    ('Direct LLM',      best_file('direct')),
    ('LangChain RAG',   best_file('langchain')),
    ('PaperQA2',        best_file('paperqa2')),
    ('LightRAG',        best_file('lightrag')),
    ('LlamaIndex',      best_file('llamaindex')),
    ('Haystack',        best_file('haystack')),
    ('FastGraphRAG',    best_file('fastgraphrag')),
    ('VelociRAG',       best_file('velocirag')),
    ('R2R',             best_file('r2r')),
    ('ReAct Agent',     best_file('agent_react')),
    ('Codex Sim Agent', best_file('agent_codex')),
    ('Codex Real Agent', best_file('agent_codex_real_gpt-5_5')),
    ('Claude Code Agent', best_file('agent_claude_real_opus')),
    ('ChemEvoRAG',      CHEMEVORAG),
]

# ---- Styles ----
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
data_font = Font(name='Calibri', size=11)
chem_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')  # green for ChemEvoRAG
wrap_align = Alignment(wrap_text=True, vertical='top')
center_align = Alignment(horizontal='center', vertical='top')
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin'))

wb = openpyxl.Workbook()

# ============================
# Sheet 1: Detailed QA
# ============================
ws_qa = wb.active
ws_qa.title = 'Detailed QA'

# Headers
headers = ['#', 'Paper', 'Q ID', 'Intent', 'Question', 'Ground Truth',
           'Baseline', 'System Answer', 'Score', 'Entities Found',
           'Entities Missing', 'Judge Reasoning']
for c, h in enumerate(headers, 1):
    cell = ws_qa.cell(row=1, column=c, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = center_align
    cell.border = thin_border

# Data rows
row_idx = 2
seq = 0
for baseline_name, filepath in BASELINES:
    if filepath is None or not os.path.exists(filepath):
        continue
    results = load_results(filepath)
    for r in results:
        seq += 1
        ws_qa.cell(row=row_idx, column=1, value=seq).font = data_font
        ws_qa.cell(row=row_idx, column=2, value=r.get('paper_id', '')).font = data_font
        ws_qa.cell(row=row_idx, column=3, value=r.get('question_id', '')).font = data_font
        ws_qa.cell(row=row_idx, column=4, value=r.get('intent', '')).font = data_font
        ws_qa.cell(row=row_idx, column=5, value=r.get('question', '')).font = data_font
        ws_qa.cell(row=row_idx, column=6, value=r.get('ground_truth', '') or r.get('ground_truth_answer', '')).font = data_font
        ws_qa.cell(row=row_idx, column=7, value=baseline_name).font = data_font
        ans = r.get('system_answer', '')
        if ans is None: ans = ''
        ws_qa.cell(row=row_idx, column=8, value=str(ans)).font = data_font
        ws_qa.cell(row=row_idx, column=9, value=r.get('judge_score')).font = data_font
        ef = r.get('entities_found', [])
        if ef is None: ef = []
        ws_qa.cell(row=row_idx, column=10, value=', '.join(ef) if ef else None).font = data_font
        em = r.get('entities_missing', [])
        if em is None: em = []
        ws_qa.cell(row=row_idx, column=11, value=', '.join(em) if em else None).font = data_font
        jr = r.get('judge_reasoning', '')
        ws_qa.cell(row=row_idx, column=12, value=str(jr) if jr else None).font = data_font

        # Apply borders and wrap
        for c in range(1, 13):
            cell = ws_qa.cell(row=row_idx, column=c)
            cell.border = thin_border
            if c in (5, 6, 8, 12):
                cell.alignment = wrap_align
            elif c in (1, 3, 4, 7, 9):
                cell.alignment = center_align

        # Green highlight for ChemEvoRAG
        if baseline_name == 'ChemEvoRAG':
            for c in range(1, 13):
                ws_qa.cell(row=row_idx, column=c).fill = chem_fill

        row_idx += 1

# Column widths
col_widths = {1: 6, 2: 10, 3: 8, 4: 18, 5: 50, 6: 40, 7: 16, 8: 60, 9: 8, 10: 30, 11: 30, 12: 50}
for c, w in col_widths.items():
    ws_qa.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w

ws_qa.freeze_panes = 'A2'
ws_qa.auto_filter.ref = f'A1:L{row_idx - 1}'

# ============================
# Sheet 2: Summary
# ============================
ws_sum = wb.create_sheet('Summary')

# Headers
sum_headers = ['Baseline'] + [INTENT_SHORT[i] for i in INTENTS] + ['OVERALL']
for c, h in enumerate(sum_headers, 1):
    cell = ws_sum.cell(row=1, column=c, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = center_align
    cell.border = thin_border

# Data rows
sum_row = 2
for baseline_name, filepath in BASELINES:
    if filepath is None or not os.path.exists(filepath):
        continue
    results = load_results(filepath)
    # Group by intent
    bi = defaultdict(list)
    for r in results:
        bi[r.get('intent', '?')].append(r.get('judge_score', 0))

    ws_sum.cell(row=sum_row, column=1, value=baseline_name).font = data_font
    ws_sum.cell(row=sum_row, column=1).border = thin_border

    for c, intent in enumerate(INTENTS, 2):
        scores = bi.get(intent, [])
        avg = round(sum(scores) / len(scores), 3) if scores else 0
        cell = ws_sum.cell(row=sum_row, column=c, value=avg)
        cell.font = data_font
        cell.number_format = '0.000'
        cell.alignment = center_align
        cell.border = thin_border

    all_scores = [r.get('judge_score', 0) for r in results]
    overall = round(sum(all_scores) / len(all_scores), 3) if all_scores else 0
    cell = ws_sum.cell(row=sum_row, column=7, value=overall)
    cell.font = data_font
    cell.number_format = '0.000'
    cell.alignment = center_align
    cell.border = thin_border

    # Green highlight
    if baseline_name == 'ChemEvoRAG':
        for c in range(1, 8):
            ws_sum.cell(row=sum_row, column=c).fill = chem_fill

    sum_row += 1

# Column widths for summary
for c in range(1, 8):
    ws_sum.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 16

OUT = f'{DIR}/comparison_final.xlsx'
wb.save(OUT)
print(f'Saved: {OUT}')
print(f'QA rows: {row_idx - 2}, Summary rows: {sum_row - 2}')

# Print summary
print(f'\n{"Baseline":16s} {"ALIAS":>7s} {"ENTITY":>7s} {"PROP":>7s} {"COMP":>7s} {"COND":>7s} {"OVERALL":>7s}')
print('-' * 65)
for baseline_name, filepath in BASELINES:
    if filepath is None or not os.path.exists(filepath):
        continue
    results = load_results(filepath)
    bi = defaultdict(list)
    for r in results:
        bi[r.get('intent', '?')].append(r.get('judge_score', 0))
    vals = []
    for intent in INTENTS:
        scores = bi.get(intent, [])
        vals.append(f'{sum(scores)/len(scores):7.3f}' if scores else '    N/A')
    all_scores = [r.get('judge_score', 0) for r in results]
    print(f'{baseline_name:16s} {" ".join(vals)} {sum(all_scores)/len(all_scores):7.3f}')
